import os
from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'maxgroup-rh-2026-secure')

# Suporte local (SQLite) e nuvem (PostgreSQL)
_db_url = os.environ.get('DATABASE_URL', '')
if _db_url.startswith('postgres://'):
    _db_url = _db_url.replace('postgres://', 'postgresql://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = _db_url or 'sqlite:///maxgroup.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

LOCALIZACOES = ['315 Asa Norte', '105 Asa Sul', '103 Asa Sul', '103 Sudoeste', '103 Asa Norte']
EMPRESAS_INICIAIS = ['MaxGroup', 'MaxForm', 'MaxFoods', 'MaxFoods Restaurante', 'MaxPharma']


class Empresa(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nome = db.Column(db.String(100), nullable=False)
    cnpj = db.Column(db.String(20))
    logradouro = db.Column(db.String(200))
    numero = db.Column(db.String(10))
    complemento = db.Column(db.String(100))
    bairro = db.Column(db.String(100))
    cidade = db.Column(db.String(100))
    estado = db.Column(db.String(2))
    cep = db.Column(db.String(10))
    telefone = db.Column(db.String(20))
    email = db.Column(db.String(100))

    @property
    def endereco_completo(self):
        parts = []
        if self.logradouro:
            parts.append(self.logradouro + (f', {self.numero}' if self.numero else ''))
        if self.complemento:
            parts.append(self.complemento)
        if self.bairro:
            parts.append(self.bairro)
        if self.cidade:
            parts.append(self.cidade + (f'/{self.estado}' if self.estado else ''))
        if self.cep:
            parts.append(f'CEP {self.cep}')
        return ' — '.join(parts) if parts else ''


class Colaborador(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nome_completo = db.Column(db.String(200), nullable=False)
    empresa = db.Column(db.String(100), nullable=False)
    localizacao = db.Column(db.String(100), nullable=False)
    tipo_contrato = db.Column(db.String(20))   # CLT, PJ, Estágio, Terceirizado
    cargo = db.Column(db.String(100))
    horario = db.Column(db.String(200))
    email = db.Column(db.String(100))
    data_admissao = db.Column(db.Date)
    data_aniversario = db.Column(db.Date)
    contato = db.Column(db.String(50))
    rg = db.Column(db.String(30))
    cpf = db.Column(db.String(20))
    data_expedicao_rg = db.Column(db.Date)
    local_expedicao_rg = db.Column(db.String(100))
    nacionalidade = db.Column(db.String(100))
    naturalidade = db.Column(db.String(100))
    nome_pai = db.Column(db.String(200))
    nome_mae = db.Column(db.String(200))
    numero_ctps = db.Column(db.String(50))
    data_emissao_ctps = db.Column(db.Date)
    numero_agencia = db.Column(db.String(20))
    numero_conta = db.Column(db.String(30))
    remuneracao = db.Column(db.Float, default=0.0)
    premiacao = db.Column(db.Float, default=0.0)
    vale_transporte = db.Column(db.Float, default=0.0)
    auxilio_transporte = db.Column(db.Float, default=0.0)
    vale_alimentacao = db.Column(db.Float, default=0.0)
    assiduidade = db.Column(db.Float, default=0.0)
    comissao = db.Column(db.Float, default=0.0)
    data_desligamento = db.Column(db.Date, nullable=True)
    # Endereço residencial
    end_cep = db.Column(db.String(10))
    end_logradouro = db.Column(db.String(200))
    end_numero = db.Column(db.String(10))
    end_complemento = db.Column(db.String(100))
    end_bairro = db.Column(db.String(100))
    end_cidade = db.Column(db.String(100))
    end_estado = db.Column(db.String(2))
    criado_em = db.Column(db.DateTime, default=datetime.utcnow)

    @property
    def ativo(self):
        return self.data_desligamento is None

    @property
    def total_proventos(self):
        return sum(filter(None, [
            self.remuneracao, self.premiacao, self.vale_transporte,
            self.auxilio_transporte, self.vale_alimentacao,
            self.assiduidade, self.comissao
        ]))

    @property
    def endereco_completo(self):
        parts = []
        if self.end_logradouro:
            parts.append(self.end_logradouro + (f', {self.end_numero}' if self.end_numero else ''))
        if self.end_complemento:
            parts.append(self.end_complemento)
        if self.end_bairro:
            parts.append(self.end_bairro)
        if self.end_cidade:
            parts.append(self.end_cidade + (f'/{self.end_estado}' if self.end_estado else ''))
        if self.end_cep:
            parts.append(f'CEP {self.end_cep}')
        return ' — '.join(parts) if parts else ''


def parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except Exception:
        return None


def fmt_date(d):
    return d.strftime('%d/%m/%Y') if d else ''


def fmt_currency(value):
    if value is None:
        return 'R$ 0,00'
    return 'R$ {:,.2f}'.format(value).replace(',', 'X').replace('.', ',').replace('X', '.')


app.jinja_env.filters['fmt_date'] = fmt_date
app.jinja_env.filters['fmt_currency'] = fmt_currency


@app.context_processor
def inject_globals():
    empresas_db = [e.nome for e in Empresa.query.order_by(Empresa.nome).all()]
    return {
        'now': datetime.now(),
        'empresas': empresas_db,
        'localizacoes': LOCALIZACOES,
    }


def migrate_db():
    """Adiciona colunas novas a tabelas existentes sem perder dados."""
    new_cols = [
        ('colaborador', 'end_cep', 'VARCHAR(10)'),
        ('colaborador', 'end_logradouro', 'VARCHAR(200)'),
        ('colaborador', 'end_numero', 'VARCHAR(10)'),
        ('colaborador', 'end_complemento', 'VARCHAR(100)'),
        ('colaborador', 'end_bairro', 'VARCHAR(100)'),
        ('colaborador', 'end_cidade', 'VARCHAR(100)'),
        ('colaborador', 'end_estado', 'VARCHAR(2)'),
        ('colaborador', 'tipo_contrato', 'VARCHAR(20)'),
        ('colaborador', 'cargo', 'VARCHAR(100)'),
        ('colaborador', 'horario', 'VARCHAR(200)'),
        ('colaborador', 'email', 'VARCHAR(100)'),
    ]
    from sqlalchemy import text
    is_pg = 'postgresql' in app.config['SQLALCHEMY_DATABASE_URI']
    for table, col, col_type in new_cols:
        try:
            if is_pg:
                # cada ALTER roda em sua própria conexão com autocommit
                with db.engine.connect().execution_options(isolation_level='AUTOCOMMIT') as conn:
                    conn.execute(text(
                        f'ALTER TABLE {table} ADD COLUMN IF NOT EXISTS {col} {col_type}'))
            else:
                with db.engine.connect() as conn:
                    conn.execute(text(f'ALTER TABLE {table} ADD COLUMN {col} {col_type}'))
                    conn.commit()
        except Exception:
            pass


# ── Colaboradores ──────────────────────────────────────────────────────────────

@app.route('/')
def index():
    empresa = request.args.get('empresa', '')
    localizacao = request.args.get('localizacao', '')
    status = request.args.get('status', 'ativo')
    search = request.args.get('search', '')

    query = Colaborador.query
    if empresa:
        query = query.filter(Colaborador.empresa == empresa)
    if localizacao:
        query = query.filter(Colaborador.localizacao == localizacao)
    if search:
        query = query.filter(Colaborador.nome_completo.ilike(f'%{search}%'))
    if status == 'ativo':
        query = query.filter(Colaborador.data_desligamento == None)
    elif status == 'desligado':
        query = query.filter(Colaborador.data_desligamento != None)

    colaboradores = query.order_by(Colaborador.nome_completo).all()
    total_ativos = Colaborador.query.filter(Colaborador.data_desligamento == None).count()
    total_desligados = Colaborador.query.filter(Colaborador.data_desligamento != None).count()

    return render_template('index.html',
        colaboradores=colaboradores,
        filtro_empresa=empresa, filtro_localizacao=localizacao,
        filtro_status=status, filtro_search=search,
        total_ativos=total_ativos, total_desligados=total_desligados,
        total_geral=total_ativos + total_desligados,
    )


def _form_to_model(form, obj=None):
    if obj is None:
        obj = Colaborador()
    obj.nome_completo = form['nome_completo']
    obj.empresa = form['empresa']
    obj.localizacao = form['localizacao']
    obj.tipo_contrato = form.get('tipo_contrato') or None
    obj.cargo = form.get('cargo') or None
    obj.horario = form.get('horario') or None
    obj.email = form.get('email') or None
    obj.data_admissao = parse_date(form.get('data_admissao'))
    obj.data_aniversario = parse_date(form.get('data_aniversario'))
    obj.contato = form.get('contato') or None
    obj.rg = form.get('rg') or None
    obj.cpf = form.get('cpf') or None
    obj.data_expedicao_rg = parse_date(form.get('data_expedicao_rg'))
    obj.local_expedicao_rg = form.get('local_expedicao_rg') or None
    obj.nacionalidade = form.get('nacionalidade') or None
    obj.naturalidade = form.get('naturalidade') or None
    obj.nome_pai = form.get('nome_pai') or None
    obj.nome_mae = form.get('nome_mae') or None
    obj.numero_ctps = form.get('numero_ctps') or None
    obj.data_emissao_ctps = parse_date(form.get('data_emissao_ctps'))
    obj.numero_agencia = form.get('numero_agencia') or None
    obj.numero_conta = form.get('numero_conta') or None
    obj.remuneracao = float(form.get('remuneracao') or 0)
    obj.premiacao = float(form.get('premiacao') or 0)
    obj.vale_transporte = float(form.get('vale_transporte') or 0)
    obj.auxilio_transporte = float(form.get('auxilio_transporte') or 0)
    obj.vale_alimentacao = float(form.get('vale_alimentacao') or 0)
    obj.assiduidade = float(form.get('assiduidade') or 0)
    obj.comissao = float(form.get('comissao') or 0)
    obj.data_desligamento = parse_date(form.get('data_desligamento'))
    obj.end_cep = form.get('end_cep') or None
    obj.end_logradouro = form.get('end_logradouro') or None
    obj.end_numero = form.get('end_numero') or None
    obj.end_complemento = form.get('end_complemento') or None
    obj.end_bairro = form.get('end_bairro') or None
    obj.end_cidade = form.get('end_cidade') or None
    obj.end_estado = form.get('end_estado') or None
    return obj


@app.route('/colaborador/novo', methods=['GET', 'POST'])
def novo_colaborador():
    if request.method == 'POST':
        col = _form_to_model(request.form)
        db.session.add(col)
        db.session.commit()
        flash('Colaborador cadastrado com sucesso!', 'success')
        return redirect(url_for('index'))
    return render_template('form.html', colaborador=None, title='Novo Colaborador')


@app.route('/colaborador/<int:id>')
def ver_colaborador(id):
    return render_template('view.html', colaborador=Colaborador.query.get_or_404(id))


@app.route('/colaborador/<int:id>/editar', methods=['GET', 'POST'])
def editar_colaborador(id):
    colaborador = Colaborador.query.get_or_404(id)
    if request.method == 'POST':
        _form_to_model(request.form, colaborador)
        db.session.commit()
        flash('Cadastro atualizado com sucesso!', 'success')
        return redirect(url_for('ver_colaborador', id=id))
    return render_template('form.html', colaborador=colaborador, title='Editar Colaborador')


@app.route('/colaborador/<int:id>/excluir', methods=['POST'])
def excluir_colaborador(id):
    db.session.delete(Colaborador.query.get_or_404(id))
    db.session.commit()
    flash('Colaborador excluído.', 'success')
    return redirect(url_for('index'))


# ── Empresas ───────────────────────────────────────────────────────────────────

@app.route('/empresas')
def listar_empresas():
    return render_template('empresas.html', empresas=Empresa.query.order_by(Empresa.nome).all())


@app.route('/empresa/nova', methods=['GET', 'POST'])
def nova_empresa():
    if request.method == 'POST':
        emp = _form_to_empresa(request.form)
        db.session.add(emp)
        db.session.commit()
        flash(f'Empresa "{emp.nome}" cadastrada!', 'success')
        return redirect(url_for('listar_empresas'))
    return render_template('empresa_form.html', empresa=None, title='Nova Empresa')


@app.route('/empresa/<int:id>/editar', methods=['GET', 'POST'])
def editar_empresa(id):
    empresa = Empresa.query.get_or_404(id)
    if request.method == 'POST':
        _form_to_empresa(request.form, empresa)
        db.session.commit()
        flash(f'Empresa "{empresa.nome}" atualizada!', 'success')
        return redirect(url_for('listar_empresas'))
    return render_template('empresa_form.html', empresa=empresa, title='Editar Empresa')


@app.route('/empresa/<int:id>/excluir', methods=['POST'])
def excluir_empresa(id):
    empresa = Empresa.query.get_or_404(id)
    nome = empresa.nome
    db.session.delete(empresa)
    db.session.commit()
    flash(f'Empresa "{nome}" excluída.', 'success')
    return redirect(url_for('listar_empresas'))


def _form_to_empresa(form, obj=None):
    if obj is None:
        obj = Empresa()
    obj.nome = form['nome']
    obj.cnpj = form.get('cnpj') or None
    obj.logradouro = form.get('logradouro') or None
    obj.numero = form.get('numero') or None
    obj.complemento = form.get('complemento') or None
    obj.bairro = form.get('bairro') or None
    obj.cidade = form.get('cidade') or None
    obj.estado = form.get('estado') or None
    obj.cep = form.get('cep') or None
    obj.telefone = form.get('telefone') or None
    obj.email = form.get('email') or None
    return obj


# ── FOPAG ──────────────────────────────────────────────────────────────────────

@app.route('/exportar-fopag')
def exportar_fopag():
    empresa_filtro = request.args.get('empresa', '')
    localizacao = request.args.get('localizacao', '')
    mes_ref = request.args.get('mes_ref', datetime.now().strftime('%Y-%m'))

    query = Colaborador.query.filter(Colaborador.data_desligamento == None)
    if empresa_filtro:
        query = query.filter(Colaborador.empresa == empresa_filtro)
    if localizacao:
        query = query.filter(Colaborador.localizacao == localizacao)

    colaboradores = query.order_by(Colaborador.empresa, Colaborador.nome_completo).all()

    try:
        mes_nome = datetime.strptime(mes_ref + '-01', '%Y-%m-%d').strftime('%B %Y').upper()
    except Exception:
        mes_nome = mes_ref

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'FOPAG'

    PRETO = '141414'
    OURO = 'C9A84C'

    def borda():
        s = Side(style='thin', color='D0C8B0')
        return Border(left=s, right=s, top=s, bottom=s)

    ws.merge_cells('A1:S1')
    t = ws['A1']
    t.value = 'MAXGROUP — FOLHA DE PAGAMENTO'
    t.font = Font(name='Calibri', bold=True, size=18, color=OURO)
    t.fill = PatternFill(start_color=PRETO, end_color=PRETO, fill_type='solid')
    t.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 38

    ws.merge_cells('A2:S2')
    t2 = ws['A2']
    t2.value = f'Competência: {mes_nome}' + (f'   |   {empresa_filtro}' if empresa_filtro else '') + (f'   |   {localizacao}' if localizacao else '')
    t2.font = Font(name='Calibri', size=11, color='888888')
    t2.fill = PatternFill(start_color='1E1E1E', end_color='1E1E1E', fill_type='solid')
    t2.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 6

    HDR_ROW = 4
    ws.row_dimensions[HDR_ROW].height = 40

    columns = [
        ('Nº', 4), ('NOME COMPLETO', 36), ('EMPRESA', 22), ('LOCALIZAÇÃO', 18),
        ('ADMISSÃO', 13), ('CPF', 16), ('RG', 14), ('CTPS', 14), ('CONTATO', 16),
        ('AGÊNCIA', 10), ('CONTA', 14),
        ('REMUNERAÇÃO', 16), ('PREMIAÇÃO', 14), ('V. TRANSPORTE', 15),
        ('AUX. TRANSPORTE', 17), ('V. ALIMENTAÇÃO', 15),
        ('ASSIDUIDADE', 14), ('COMISSÃO', 13), ('TOTAL', 16),
    ]

    for ci, (label, width) in enumerate(columns, 1):
        c = ws.cell(row=HDR_ROW, column=ci, value=label)
        c.font = Font(name='Calibri', bold=True, color=OURO, size=10)
        c.fill = PatternFill(start_color=PRETO, end_color=PRETO, fill_type='solid')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = borda()
        ws.column_dimensions[get_column_letter(ci)].width = width

    CURR_FMT = '"R$"#,##0.00'

    for i, col in enumerate(colaboradores, 1):
        row = HDR_ROW + i
        ws.row_dimensions[row].height = 17
        fill_color = 'FAFAF7' if i % 2 == 0 else 'FFFFFF'
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
        base_font = Font(name='Calibri', size=10)

        values = [
            i, col.nome_completo, col.empresa, col.localizacao,
            col.data_admissao, col.cpf, col.rg, col.numero_ctps, col.contato,
            col.numero_agencia, col.numero_conta,
            col.remuneracao or 0, col.premiacao or 0, col.vale_transporte or 0,
            col.auxilio_transporte or 0, col.vale_alimentacao or 0,
            col.assiduidade or 0, col.comissao or 0, col.total_proventos,
        ]

        for ci, val in enumerate(values, 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.font = base_font
            c.fill = row_fill
            c.border = borda()
            if ci == 1:
                c.alignment = Alignment(horizontal='center', vertical='center')
            elif ci == 5 and isinstance(val, date):
                c.number_format = 'DD/MM/YYYY'
                c.alignment = Alignment(horizontal='center', vertical='center')
            elif ci >= 12:
                c.number_format = CURR_FMT
                c.alignment = Alignment(horizontal='right', vertical='center')
            else:
                c.alignment = Alignment(horizontal='left', vertical='center')

    if colaboradores:
        tot_row = HDR_ROW + len(colaboradores) + 1
        ws.row_dimensions[tot_row].height = 22
        ws.merge_cells(f'A{tot_row}:K{tot_row}')
        lbl = ws.cell(row=tot_row, column=1, value=f'TOTAL  ({len(colaboradores)} colaboradores ativos)')
        lbl.font = Font(name='Calibri', bold=True, color=OURO, size=11)
        lbl.fill = PatternFill(start_color=PRETO, end_color=PRETO, fill_type='solid')
        lbl.alignment = Alignment(horizontal='center', vertical='center')
        lbl.border = borda()
        for ci in range(12, 20):
            col_ltr = get_column_letter(ci)
            c = ws.cell(row=tot_row, column=ci,
                        value=f'=SUM({col_ltr}{HDR_ROW+1}:{col_ltr}{HDR_ROW+len(colaboradores)})')
            c.font = Font(name='Calibri', bold=True, color=OURO, size=11)
            c.fill = PatternFill(start_color=PRETO, end_color=PRETO, fill_type='solid')
            c.number_format = CURR_FMT
            c.alignment = Alignment(horizontal='right', vertical='center')
            c.border = borda()

    ws.freeze_panes = f'A{HDR_ROW + 1}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f'FOPAG_{mes_nome.replace(" ", "-")}'
    if empresa_filtro:
        filename += f'_{empresa_filtro.replace(" ", "_")}'
    filename += '.xlsx'

    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        migrate_db()
        if Empresa.query.count() == 0:
            for nome in EMPRESAS_INICIAIS:
                db.session.add(Empresa(nome=nome))
            db.session.commit()
    port = int(os.environ.get('PORT', 5000))
    debug = os.environ.get('FLASK_ENV') == 'development'
    print(f'\n  MaxGroup RH — http://localhost:{port}\n')
    app.run(host='0.0.0.0', port=port, debug=debug)
